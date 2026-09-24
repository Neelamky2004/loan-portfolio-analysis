import os
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
DB_URL = os.getenv("DATABASE_URL", "postgresql://postgres@localhost:5432/loans")

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TILE_FILL = PatternFill("solid", fgColor="EEF2F8")
THIN = Side(style="thin", color="C9D1E0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RISK_MAP = [("A", "Low"), ("B", "Low"), ("C", "Medium"), ("D", "Medium"), ("E", "High"), ("F", "High"), ("G", "High")]

SQL_SHEETS = {
    "Monthly_Trend": "select * from v_monthly_disbursal",
    "Status_Buckets": "select * from v_status_buckets",
    "Vintage": "select * from v_vintage",
    "Purpose": "select * from v_purpose_performance",
    "States": "select * from v_state_summary",
    "Collections_List": "select * from v_collections_list",
}

LOAN_COLS = [
    "loan_id", "issue_date", "grade", "purpose", "state", "funded_amount", "int_rate",
    "term_months", "loan_status", "outstanding_principal", "principal_received",
]


def write_table(ws, df, start_row=1):
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
    for j, col in enumerate(df.columns, 1):
        width = max(len(str(col)), *(len(str(v)) for v in df.iloc[:200, j - 1])) if len(df) else len(col)
        ws.column_dimensions[get_column_letter(j)].width = min(width + 3, 28)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def money_format(ws, headers, fmt="#,##0"):
    for j, col in enumerate(headers, 1):
        if any(k in col for k in ("amount", "principal", "funded", "lost", "installment")):
            for cell in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell[0].number_format = fmt


def build_loan_data(wb, loans):
    ws = wb.create_sheet("Loan_Data")
    df = loans[LOAN_COLS].copy()
    write_table(ws, df)
    band_col = len(LOAN_COLS) + 1
    h = ws.cell(row=1, column=band_col, value="risk_band")
    h.fill, h.font = HEADER_FILL, HEADER_FONT
    for r in range(2, len(df) + 2):
        ws.cell(row=r, column=band_col, value=f"=VLOOKUP(C{r},Risk_Map!$A$2:$B$8,2,FALSE)")
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    ws.column_dimensions[get_column_letter(band_col)].width = 12
    ws.auto_filter.ref = f"A1:{get_column_letter(band_col)}{len(df) + 1}"
    return len(df) + 1


def build_risk_map(wb):
    ws = wb.create_sheet("Risk_Map")
    write_table(ws, pd.DataFrame(RISK_MAP, columns=["grade", "risk_band"]))


def build_grade_summary(wb, last):
    ws = wb.create_sheet("Grade_Summary")
    headers = ["grade", "risk_band", "loans", "funded_amount", "avg_int_rate", "closed_loans",
               "defaulted_loans", "default_rate", "outstanding_principal"]
    write_table(ws, pd.DataFrame(columns=headers))
    g = f"Loan_Data!$C$2:$C${last}"
    s = f"Loan_Data!$I$2:$I${last}"
    for i, (grade, _) in enumerate(RISK_MAP, 2):
        ws[f"A{i}"] = grade
        ws[f"B{i}"] = f"=VLOOKUP(A{i},Risk_Map!$A$2:$B$8,2,FALSE)"
        ws[f"C{i}"] = f"=COUNTIFS({g},A{i})"
        ws[f"D{i}"] = f"=SUMIFS(Loan_Data!$F$2:$F${last},{g},A{i})"
        ws[f"E{i}"] = f"=AVERAGEIFS(Loan_Data!$G$2:$G${last},{g},A{i})"
        ws[f"F{i}"] = (f'=COUNTIFS({g},A{i},{s},"Fully Paid")+COUNTIFS({g},A{i},{s},"Charged Off")'
                       f'+COUNTIFS({g},A{i},{s},"Default")')
        ws[f"G{i}"] = f'=COUNTIFS({g},A{i},{s},"Charged Off")+COUNTIFS({g},A{i},{s},"Default")'
        ws[f"H{i}"] = f"=IFERROR(G{i}/F{i},0)"
        ws[f"I{i}"] = f"=SUMIFS(Loan_Data!$J$2:$J${last},{g},A{i})"
        ws[f"D{i}"].number_format = ws[f"I{i}"].number_format = "#,##0"
        ws[f"E{i}"].number_format = "0.00"
        ws[f"H{i}"].number_format = "0.0%"
    t = len(RISK_MAP) + 2
    ws[f"A{t}"] = "Total"
    ws[f"A{t}"].font = Font(bold=True)
    for col in "CDFGI":
        ws[f"{col}{t}"] = f"=SUM({col}2:{col}{t - 1})"
        ws[f"{col}{t}"].font = Font(bold=True)
        ws[f"{col}{t}"].number_format = "#,##0"
    ws[f"H{t}"] = f"=G{t}/F{t}"
    ws[f"H{t}"].number_format = "0.0%"
    ws.conditional_formatting.add(
        f"H2:H{t - 1}",
        ColorScaleRule(start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50,
                       mid_color="FFEB9C", end_type="max", end_color="F8CBAD"),
    )
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 18
    return ws


def build_dashboard(wb, last):
    ws = wb["Dashboard"]
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Loan Portfolio Dashboard"
    ws["B2"].font = Font(bold=True, size=18, color=NAVY)
    ws["B3"] = "LendingClub loans issued 2007-2011"
    ws["B3"].font = Font(italic=True, color="666666")

    s = f"Loan_Data!$I$2:$I${last}"
    tiles = [
        ("Total Loans", f"=COUNTA(Loan_Data!$A$2:$A${last})", "#,##0"),
        ("Total Funded ($)", f"=SUM(Loan_Data!$F$2:$F${last})", "#,##0"),
        ("Avg Interest Rate", f"=AVERAGE(Loan_Data!$G$2:$G${last})/100", "0.00%"),
        ("Default Rate (closed loans)",
         f'=(COUNTIFS({s},"Charged Off")+COUNTIFS({s},"Default"))/(COUNTIFS({s},"Fully Paid")'
         f'+COUNTIFS({s},"Charged Off")+COUNTIFS({s},"Default"))', "0.00%"),
        ("Outstanding Principal ($)", f"=SUM(Loan_Data!$J$2:$J${last})", "#,##0"),
        ("Accounts in Collections", f"=COUNTA(Collections_List!$A:$A)-1", "#,##0"),
    ]
    cols = ["B", "D", "F", "H", "J", "L"]
    for col, (label, formula, fmt) in zip(cols, tiles):
        nxt = get_column_letter(ws[f"{col}5"].column + 1)
        ws.merge_cells(f"{col}5:{nxt}5")
        ws.merge_cells(f"{col}6:{nxt}6")
        ws[f"{col}5"] = label
        ws[f"{col}5"].font = Font(size=9, bold=True, color="555555")
        ws[f"{col}6"] = formula
        ws[f"{col}6"].font = Font(size=16, bold=True, color=NAVY)
        ws[f"{col}6"].number_format = fmt
        for r in (5, 6):
            for c in (col, nxt):
                ws[f"{c}{r}"].fill = TILE_FILL
                ws[f"{c}{r}"].border = BOX
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[6].height = 30
    for c in range(1, 15):
        ws.column_dimensions[get_column_letter(c)].width = 11

    mt = wb["Monthly_Trend"]
    n = mt.max_row
    line = LineChart()
    line.title = "Monthly Funded Amount"
    line.y_axis.title = "Funded ($)"
    line.height, line.width = 7.5, 16
    line.add_data(Reference(mt, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    line.set_categories(Reference(mt, min_col=1, min_row=2, max_row=n))
    line.legend = None
    ws.add_chart(line, "B9")

    gs = wb["Grade_Summary"]
    bar = BarChart()
    bar.title = "Default Rate by Grade"
    bar.height, bar.width = 7.5, 12
    bar.add_data(Reference(gs, min_col=8, min_row=1, max_row=8), titles_from_data=True)
    bar.set_categories(Reference(gs, min_col=1, min_row=2, max_row=8))
    bar.y_axis.numFmt = "0%"
    bar.legend = None
    ws.add_chart(bar, "J9")

    sb = wb["Status_Buckets"]
    bar2 = BarChart()
    bar2.type = "bar"
    bar2.title = "Loans by Status Bucket"
    bar2.height, bar2.width = 7.5, 16
    bar2.add_data(Reference(sb, min_col=2, min_row=1, max_row=sb.max_row), titles_from_data=True)
    bar2.set_categories(Reference(sb, min_col=1, min_row=2, max_row=sb.max_row))
    bar2.legend = None
    ws.add_chart(bar2, "B25")

    vt = wb["Vintage"]
    line2 = LineChart()
    line2.title = "Default Rate by Issue Quarter (%)"
    line2.height, line2.width = 7.5, 12
    line2.add_data(Reference(vt, min_col=6, min_row=1, max_row=vt.max_row), titles_from_data=True)
    line2.set_categories(Reference(vt, min_col=2, min_row=2, max_row=vt.max_row))
    line2.legend = None
    ws.add_chart(line2, "J25")


def main():
    OUT.mkdir(exist_ok=True)
    engine = create_engine(DB_URL)
    with engine.connect() as conn:
        loans = pd.read_sql(text("select * from loans order by issue_date, loan_id"), conn)
        tables = {name: pd.read_sql(text(q), conn) for name, q in SQL_SHEETS.items()}
        checks = pd.read_sql(text((ROOT / "sql" / "02_quality_checks.sql").read_text()), conn)

    checks["status"] = checks["failed_rows"].map(lambda x: "PASS" if x == 0 else "REVIEW")
    for name, df in tables.items():
        for col in df.columns:
            if df[col].dtype == object and col.endswith(("month", "date")):
                df[col] = pd.to_datetime(df[col])
    tables["Vintage"]["issue_quarter"] = (
        tables["Vintage"]["issue_year"].astype(str) + "-" + tables["Vintage"]["issue_quarter"]
    )

    wb = Workbook()
    wb.active.title = "Dashboard"
    for name, df in tables.items():
        ws = wb.create_sheet(name)
        write_table(ws, df)
        money_format(ws, list(df.columns))
    wb["Monthly_Trend"].column_dimensions["A"].width = 14
    for row in wb["Monthly_Trend"].iter_rows(min_row=2, max_col=1):
        row[0].number_format = "mmm-yyyy"

    cl = wb["Collections_List"]
    n = cl.max_row
    red = PatternFill("solid", fgColor="F8CBAD")
    amber = PatternFill("solid", fgColor="FFEB9C")
    cl.conditional_formatting.add(f"E2:E{n}", CellIsRule(operator="equal", formula=['"High"'], fill=red))
    cl.conditional_formatting.add(f"E2:E{n}", CellIsRule(operator="equal", formula=['"Medium"'], fill=amber))

    dq = wb.create_sheet("Data_Checks")
    write_table(dq, checks)
    dq.conditional_formatting.add(f"C2:C{len(checks) + 1}",
                                  CellIsRule(operator="equal", formula=['"REVIEW"'], fill=amber))

    build_risk_map(wb)
    last = build_loan_data(wb, loans)
    build_grade_summary(wb, last)
    build_dashboard(wb, last)

    order = ["Dashboard", "Grade_Summary", "Monthly_Trend", "Status_Buckets", "Vintage", "Purpose",
             "States", "Collections_List", "Data_Checks", "Loan_Data", "Risk_Map"]
    wb._sheets = [wb[s] for s in order]
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT / "loan_portfolio_report.xlsx")

    loans.assign(risk_band=loans["grade"].map(dict(RISK_MAP))).to_csv(OUT / "loans_clean.csv", index=False)
    print("saved output/loan_portfolio_report.xlsx and output/loans_clean.csv")


if __name__ == "__main__":
    main()
